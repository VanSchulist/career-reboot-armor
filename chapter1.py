# -*- coding: utf-8 -*-
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import set_cell_background, set_cell_margins, add_callout

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(20)
    h.paragraph_format.space_after = Pt(10)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.bold = True
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    return h

def build_chapter_1(doc):
    add_heading_1(doc, "第一章：事实清算——背景调查（Background Check）真相与心理防御重建")
    
    add_heading_2(doc, "1. 问题：断层恐惧与被动自裁")
    doc.add_paragraph(
        "许多经历过 1 至 4 年长期家庭照护（如阿尔茨海默病长辈、严重中风亲属、重度慢性病晚期）或重大伤病康复的求职者，"
        "在准备重新打开求职软件时，最先遭遇的不是外界的拒绝，而是内心的剧烈恐慌。他们普遍持有一种隐性假设："
        "“背景调查公司像福尔摩斯一样无所不知，他们会调取我所有的银行流水、全网社保记录，一旦发现我中间有 2 年没上班，"
        "就会在背调报告上亮起红灯（Red Flag），认定我不诚实或是不良候选人。”"
    )
    doc.add_paragraph(
        "这种缺乏根据的恐慌，直接导致求职者陷入两种自毁行为：要么在投递简历前就自我否定，彻底放弃尝试中高层匹配岗位；"
        "要么在简历上铤而走险，私自推迟上一份真实全职工作的离职年份，试图把空白期抹平。"
    )

    add_heading_2(doc, "2. 为什么发生：商业调查与刑事侦查的认知错位")
    doc.add_paragraph(
        "这种恐慌源于对第三方商业背调机构（如跨国机构 HireRight、First Advantage、Sterling、Equifax The Work Number，"
        "以及国内的全景求是、太和鼎信等）运作机制的无知。求职者误以为背调是“全网主动侦查”，而实际上它只是一场严格受制于"
        "法律框架（如美国 FCRA《公平信用报告法》、中国《个人信息保护法》）的商业自动化数据比对流程。"
    )
    doc.add_paragraph(
        "第三方背调机构的核心商业逻辑是按单收费（几十元到几百元一单），其交付物是《客观数据比对表》，"
        "其根本职责是：【核实求职者在授权表上主动填写的内容是否真实】，而不是去【主动调查求职者没写的时间里干了什么】。"
        "在法律和商业规则中，“一段时间未在任何企业就业（Career Break）”属于个人合法生活状态，本身绝非任何形式的欺诈或违规。"
    )

    add_heading_2(doc, "3. 常见错误：用真正的欺诈掩盖合法的休假")
    doc.add_paragraph(
        "面对空白期，求职者最容易犯下致命错误。以下列举三种典型的高危行为："
    )
    
    # 错误对比表格
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    headers = ["常见错误做法", "求职者的侥幸心理", "背调与HR现场的真实下场"]
    col_widths = [1.8, 2.0, 2.4]
    
    for i, title in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, "EDF2F7")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    error_data = [
        ("推迟前雇员离职年份\n（例如真实2022年离职，硬写成2024年）", "“公司HR换了好几茬，只要跟前同事打好招呼，背调肯定查不出来。”", "【直接撤销Offer】背调机构发函至前公司HR总机或调取个税退工单，离职日期出现超过1个月的差异，直接标记为“严重不符（Discrepancy）”，判定求职欺诈。"),
        ("找皮包/空壳公司挂靠开假证明\n（购买虚假离职证明与假发薪流水）", "“只要有印章和电话接听，表面上天衣无缝。”", "【法律与征信暴雷】背调机构通过工商穿透、对公账户比对或社保缴纳清单直接识别空壳挂靠，不仅Offer被撤，更可能被大厂列入永久雇佣黑名单。"),
        ("在简历顶端过度倾诉家庭悲剧\n（把照护家务写满整页）", "“我要向HR展示我这几年没有闲着，我在统筹医疗和家庭支出。”", "【初筛即淘汰】HR和用人经理认为该写法极不专业，且触发雇主对“家庭琐事随时导致缺勤”的潜意识偏见，直接过滤。")
    ]
    
    for row_idx, data in enumerate(error_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, 80, 80, 100, 100)
            if row_idx % 2 == 1:
                set_cell_background(cell, "FFFFFF")
            else:
                set_cell_background(cell, "F7FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if col_idx == 2:
                run.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
            else:
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading_2(doc, "4. 正确方法：双轨申报与事实基准线锚定")
    doc.add_paragraph(
        "化解断层恐惧的核心方法，是建立极其严密的“双轨申报意识”：\n"
        "【第一轨：法定背调授权表（Legal Background Check Form）】—— 必须 100% 绝对诚实、毫厘不差。只填写你能拿出官方税单、离职证明或社保核实的历史正式单位，日期精确到月。对于照护休假期间，如实填写“职业休假（Career Break）”或留空，绝不无中生有。\n"
        "【第二轨：商业求职简历（Commercial Resume）】—— 属于市场化能力展示单，重点在于通过合规的“独立咨询、过桥项目、现代化技能”消除时间断崖（详见第二章），使简历顺利通过算法解析与人类筛选。"
    )

    add_callout(
        doc,
        "在招聘法律中，简历（Resume）是营销沟通材料，而背调授权书（Consent Form）是法定确认文件。\n"
        "只要你在背调授权表上填报的每一段正式雇佣记录均真实有效（入职时间、离职时间、正式法定头衔均有据可查），"
        "中间存在的空白期在背调合规判定上属于【标准无记录（Standard Non-Employment）】，绝不会触发合规红线！",
        title="背调红线铁律"
    )

    add_heading_2(doc, "5. 具体步骤：四步锁定个人历史数据基准线")
    doc.add_paragraph(
        "在修改简历之前，求职者必须在物理上彻底搞清楚自己留存在外部公共数据库中的档案到底记录了什么。请按以下四步操作："
    )
    doc.add_paragraph("第一步：下载个人官方就业与参保记录\n- 国内求职者：登录“掌上12333”App或本地人社/社保平台，导出《个人权益记录单》（包含历任雇主的参保起止年月）及个人所得税App纳税记录；\n- 跨国/海外求职者：访问 Equifax The Work Number 官网，免费申请并下载自己的《Employment Data Report (EDR)》，核实系统内记录的离职月份。")
    doc.add_paragraph("第二步：整理《法定历史事实基准表》\n将所有真实全职经历列为不可篡改的锚点：【公司法定名称】、【入职年月】、【离职证明上的最后离职年月】、【离职证明上的官方职务名称】。这些数据一经确定，在任何系统里都必须保持绝对一致。")
    doc.add_paragraph("第三步：圈定真实空白期边界\n计算上一段正式离职证明上的最后月份至今天的真实间隔（例如：2022年10月至2025年3月，共计29个月）。在心理上坦然承认这29个月是家庭医疗与照护的休假事实。")
    doc.add_paragraph("第四步：确立前雇主证明人渠道\n联系上家公司的直属领导或当年的HR同事，确认前公司的公开发薪座机或HR查询邮箱是否依然有效，确保未来背调机构发函时有人能够客观证明你在职期间的良好表现。")

    add_heading_2(doc, "6. 示例：标准背调与两类申报对照")
    doc.add_paragraph(
        "以求职者张女士（原某外企市场总监，2022年8月辞职全职照护病重母亲至2024年底）为例，观察其在背调环节的合规操作："
    )
    doc.add_paragraph(
        "【情境】：新公司发放 Offer，进入第三方背调环节，系统要求填写过去 5 年的履历。\n"
        "【张女士的正确做法】在背调系统填报页面，她如实填写：\n"
        "• 2017.03 – 2022.08：某某科技有限公司 | 市场总监（提供当年离职证明与个税单核验）\n"
        "• 2022.09 – 2024.12：职业休假 / 家庭照护（Caregiver Leave）\n"
        "• 2025.01 – 至今：独立市场咨询项目（提供给朋友小微企业的咨询成果与证明人电话）\n"
        "【背调公司最终出具报告】：\n"
        "“某某科技有限公司任职时间与离职证明核验一致（Clear）；期间无违法犯罪记录（Clear）；2022.09-2024.12 标注为个人家庭事务休假（Noted / Normal Gap）。综合判定：通过（Green Pass）。”"
    )

    add_heading_2(doc, "7. 检查清单：背调防雷自测表")
    doc.add_paragraph("在向任何潜在雇主提交正式背景调查表格前，逐一核对以下 5 项：")
    doc.add_paragraph("【1】我填报的每一家前雇主起止月份，与社保缴纳或离职证明上的记录误差不超过 30 天。")
    doc.add_paragraph("【2】我填报的职务名称，与前雇主离职证明上的法定职级大体一致（不出现从‘专员’直接捏造成‘总监’的虚假跨越）。")
    doc.add_paragraph("【3】我绝未在正式背调授权表上虚构任何一段不存在劳动关系的假公司。")
    doc.add_paragraph("【4】我提供的上级证明人（Reference）已提前打过招呼，对方知道我正在求职，并愿意给予客观正面的评价。")
    doc.add_paragraph("【5】我清楚知道空白期本身是合法的个人生活权利，并在背调表上坦然填写为个人/家庭休假。")

    add_heading_2(doc, "8. 读者可以立即执行的行动")
    doc.add_paragraph(
        "【立即行动】：放下手头的简历修改，今天花 20 分钟登录个人社保/个税App（海外用户访问 The Work Number），"
        "把过去真实工作的所有离职证明拍照归档到一个专门的文件夹中，建立属于你的《真实履历事实基准线》。"
        "只要这条底线清晰确凿，后文所有的过桥重构就拥有了不可动摇的合规根基。"
    )
    
    doc.add_page_break()

print("Chapter 1 module ready.")
