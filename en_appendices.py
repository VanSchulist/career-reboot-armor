# -*- coding: utf-8 -*-
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import set_cell_background, set_cell_margins
from make_english_book import add_heading_1, add_heading_2, add_callout

def build_en_appendices(doc):
    doc.add_page_break()
    add_heading_1(doc, "Appendix Toolkit: Ready-to-Deploy Execution Assets")
    
    # -------------------------------------------------------------
    # Appendix A
    # -------------------------------------------------------------
    add_heading_2(doc, "Appendix A: The Production-Ready ATS Hybrid Resume Template")
    doc.add_paragraph(
        "This template is engineered to parse flawlessly across Workday, Greenhouse, Taleo, and iCIMS. "
        "Copy the plain text structure below, replace the bracketed placeholders with your factual data, "
        "and maintain single-column formatting:"
    )
    
    table_a = doc.add_table(rows=1, cols=1)
    table_a.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_a.autofit = False
    cell_a = table_a.cell(0, 0)
    cell_a.width = Inches(6.2)
    set_cell_background(cell_a, "FAFAFA")
    set_cell_margins(cell_a, 120, 120, 140, 140)
    p_a = cell_a.paragraphs[0]
    p_a.paragraph_format.space_before = Pt(0)
    p_a.paragraph_format.space_after = Pt(2)
    
    template_text = (
        "[FIRST NAME] [LAST NAME]\n"
        "[City, State ZIP] | [Phone Number] | [Clean Professional Email] | [LinkedIn Profile URL]\n"
        "================================================================================\n\n"
        "[TARGET PROFESSIONAL TITLE IN CAPS, e.g., SENIOR DIRECTOR OF OPERATIONS]\n"
        "[Core Capability 1] | [Core Capability 2] | [Core Capability 3]\n\n"
        "EXECUTIVE SUMMARY\n"
        "Performance-driven [Professional Title] with [X]+ years of executive experience leading [Industry/Function] across [Type of Companies, e.g., enterprise B2B environments]. "
        "Proven track record of orchestrating cross-functional teams, executing complex operational turnarounds, and delivering [Key Metric, e.g., double-digit margin expansion]. "
        "Combines deep institutional domain expertise with modern digital toolchains (Notion, Airtable, GenAI-assisted workflows) to eliminate operational latency.\n\n"
        "CORE IMPACT & VALUE HIGHLIGHTS\n"
        "• Scaled [Department/Function] from [Metric A] to [Metric B], managing annual operating budgets up to $[X]M.\n"
        "• Designed and executed [Major Initiative], resulting in a [X]% reduction in [Cost/Cycle Time] across [X] business units.\n"
        "• Institutionalized agile performance metrics and automated telemetry, driving on-time SLA fulfillment from [X]% to [X]%.\n\n"
        "PROFESSIONAL EXPERIENCE\n\n"
        "INDEPENDENT ADVISORY SERVICES | [City, State]                     [MM/YYYY] – Present\n"
        "Independent Project Consultant\n"
        "Retained by growth-stage organizations and boutique enterprises to audit [Functional Area] and modernize workflows.\n"
        "• Conducted operational diagnostic for a mid-market [Industry] client, restructuring [Process X] and eliminating [Y] hours of manual reporting.\n"
        "• Built automated tracking models utilizing [Modern Tool, e.g., Airtable/Zapier], increasing pipeline velocity by [X]%.\n"
        "• Authored comprehensive Standard Operating Procedures (SOPs) adopted company-wide, ensuring regulatory compliance and audit readiness.\n\n"
        "[PREVIOUS COMPANY NAME] | [City, State]                           [MM/YYYY] – [MM/YYYY]\n"
        "[Official Executive Job Title]\n"
        "Directed [Department Scope, e.g., regional fulfillment operations across 4 facilities], managing [X] direct reports and $[X]M P&L.\n"
        "• Spearheaded [Strategic Initiative], delivering $[X]M in recurring annual cost savings within 18 months.\n"
        "• Renegotiated strategic tier-1 vendor contracts, reducing unit procurement expenditures by [X]% while securing [Favorable Term].\n"
        "• Re-engineered team talent architecture, cutting annualized turnover from [X]% to [X]% through structured performance rubrics.\n\n"
        "[EARLIER COMPANY NAME] | [City, State]                            [MM/YYYY] – [MM/YYYY]\n"
        "[Official Functional Job Title]\n"
        "Led [Function] responsible for [Key Responsibility, e.g., client onboarding and customer success metrics].\n"
        "• Accelerated client onboarding velocity by [X]%, increasing net retention rates (NDR) to [X]%.\n"
        "• Designed company-wide training curriculum that onboarded [X] incoming specialists with zero customer churn.\n\n"
        "EDUCATION & CREDENTIALS\n"
        "• Bachelor of Science / Arts in [Major] | [University Name, City, State]\n"
        "• [Industry Credential, e.g., PMP - Project Management Professional, APICS CSCP, or CPA]"
    )
    run_t = p_a.add_run(template_text)
    run_t.font.name = 'Consolas'
    run_t.font.size = Pt(8.5)
    run_t.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # -------------------------------------------------------------
    # Appendix B
    # -------------------------------------------------------------
    add_heading_2(doc, "Appendix B: The Grilling Room — 10 Tough Interview Questions & Defensive Scripts")
    doc.add_paragraph(
        "Beyond the general 30-second pivot in Chapter 5, interviewers frequently deploy pointed, skeptical probes. "
        "Use these battle-tested, word-for-word scripts to maintain absolute professional authority:"
    )

    qa_list = [
        ("Probe 1: \"Why did you personally step away to handle caregiving instead of hiring full-time professional nurses?\"",
         "The Script: \"During the acute intervention and diagnostic phase, critical medical decisions, specialized legal power-of-attorney execution, and clinical advocacy required immediate next-of-kin authority. Once the medical framework was stabilized and transferred to permanent long-term clinical care, my phase concluded. My role was executive stabilization, and that mission was fully achieved.\""),
         
        ("Probe 2: \"How have you supported yourself financially during this gap? Are you facing acute financial distress?\"",
         "The Script: \"Prior to taking planned leave, I established disciplined, multi-year emergency financial reserves. Coupled with ongoing selective project advisory work, my financial foundation is exceptionally solid. I am here today purely because your operational expansion represents the exact commercial challenge I am best equipped to solve.\""),

        ("Probe 3: \"This is a high-pressure, fast-paced environment. How do I know your stamina hasn't atrophied?\"",
         "The Script: \"Managing complex healthcare crises while simultaneously delivering advisory milestones requires an intensity of focus, crisis tolerance, and multi-threaded problem-solving that standard office environments rarely test. My operational stamina has been thoroughly tested and proven under real adversity. I thrive in high-velocity environments.\""),

        ("Probe 4: \"What happens if your family member's health relapses? Will you need to resign again?\"",
         "The Script: \"Under no circumstances. The family situation has reached a complete, permanent, and definitive resolution [or: is backed by a fully contracted, multi-tiered institutional care infrastructure]. My professional obligations have absolute, uncompromised priority.\""),

        ("Probe 5: \"This role represents a step down in title and salary compared to your previous director-level position. Will you get bored and leave?\"",
         "The Script: \"I chose to target this role with eyes wide open. For my re-entry chapter, I prioritized organizational health, immediate operational impact, and team alignment over vanity titles. I am eager to roll up my sleeves and let my baseline results speak for themselves.\""),

        ("Probe 6: \"How do I know this 'Independent Consulting' on your resume isn't just a placeholder you made up?\"",
         "The Script: \"I completely respect your diligence. I have a de-identified portfolio of client deliverables on my laptop that I am happy to walk through right now. Furthermore, my primary client stakeholder is listed on my professional reference sheet and welcomes your call during background screening.\""),

        ("Probe 7: \"If you enjoyed consulting, why return to a structured corporate W-2 role?\"",
         "The Script: \"Consulting offers agility, but external advisors only recommend—they do not own the long-term execution. I am an operator at heart. I do my best work when embedded with a dedicated team, leveraging organizational resources to build enterprise value over the long haul.\""),

        ("Probe 8: \"You had a medical hiatus. Can you legally pass our corporate physical and travel demands?\"",
         "The Script: \"Prior to launching my re-entry campaign, I completed full clinical evaluations with my medical team. I have been given full, unrestricted clearance for all full-time professional duties and regular travel. I welcome your standard occupational health screening.\""),

        ("Probe 9: \"Your direct supervisor may be significantly younger than you. Will that create friction?\"",
         "The Script: \"In modern organizations, reporting structures reflect functional accountability, not chronological age. I deeply respect technical velocity and modern market instincts in younger leaders, while bringing complementary operational stability and risk management. I welcome working alongside sharp leaders of any age.\""),

        ("Probe 10: \"What are your compensation expectations, and are you willing to accept a market reset?\"",
         "The Script: \"Based on current market parameters for this functional scope, my target range is [$X to $Y]. As a returning executive, I possess tactical flexibility on base compensation in exchange for performance-tied incentives and meaningful equity participation.\"")
    ]

    for q, a in qa_list:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        run_q = p_q.add_run(q)
        run_q.font.bold = True
        run_q.font.size = Pt(10)
        run_q.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_after = Pt(6)
        run_a = p_a.add_run(a)
        run_a.font.size = Pt(9.5)
        run_a.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # -------------------------------------------------------------
    # Appendix C
    # -------------------------------------------------------------
    add_heading_2(doc, "Appendix C: Micro-Consulting Statement of Work (SOW) & Reference Agreement")
    doc.add_paragraph(
        "To ensure your Chapter 2 bridge engagement is 100% legally and commercially defensible during background checks, "
        "execute this standardized 1-page agreement with your business contact:"
    )

    table_c = doc.add_table(rows=1, cols=1)
    table_c.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_c.autofit = False
    cell_c = table_c.cell(0, 0)
    cell_c.width = Inches(6.2)
    set_cell_background(cell_c, "FAFAFA")
    set_cell_margins(cell_c, 120, 120, 140, 140)
    p_c = cell_c.paragraphs[0]
    p_c.paragraph_format.space_before = Pt(0)
    p_c.paragraph_format.space_after = Pt(2)
    
    sow_text = (
        "INDEPENDENT ADVISORY SERVICES STATEMENT OF WORK (SOW)\n\n"
        "This Statement of Work is entered into on [Date], by and between [Client Entity / Business Name] (\"Client\"), "
        "and [Your Full Legal Name] (\"Consultant\").\n\n"
        "1. SCOPE OF ENGAGEMENT & DELIVERABLES\n"
        "Client hereby retains Consultant to perform targeted operational diagnostic and workflow optimization advisory services. "
        "Consultant shall deliver the following specific work product:\n"
        "• Diagnostic evaluation of [Target Workflow, e.g., Vendor Onboarding / Order Reconciliation];\n"
        "• Delivery of a finalized Standard Operating Procedure (SOP) / Relational Tracking Database; and\n"
        "• Executive debrief session with Client stakeholders.\n\n"
        "2. TERM & PROFESSIONAL COMPENSATION\n"
        "• Term: Services shall commence on [Start Date] and conclude on or about [End Date].\n"
        "• Consideration: Client and Consultant agree that services are rendered on an independent advisory basis for mutually agreed consideration [e.g., $1.00 nominal / $X.00 retainer / Pro-bono professional exchange].\n\n"
        "3. REFERENCE VERIFICATION & PORTFOLIO USAGE\n"
        "• Client confirms that Consultant acts as an independent contractor, not an employee.\n"
        "• Client agrees to serve as a factual, professional verification reference for Consultant upon inquiries from corporate background screening agencies, confirming dates of service and scope of engagement.\n"
        "• Consultant is authorized to showcase de-identified, non-confidential deliverables in professional career portfolios.\n\n"
        "AGREED AND ACCEPTED:\n\n"
        "For Client: ___________________________    Date: ______________\n"
        "Print Name & Title: [Client Authorized Representative]\n\n"
        "For Consultant: _______________________    Date: ______________\n"
        "Print Name: [Your Legal Name]"
    )
    run_c = p_c.add_run(sow_text)
    run_c.font.name = 'Consolas'
    run_c.font.size = Pt(8.5)
    run_c.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # -------------------------------------------------------------
    # Appendix D
    # -------------------------------------------------------------
    add_heading_2(doc, "Appendix D: The 30-Day Re-Entry Action Calendar & Accountability Checklist")
    doc.add_paragraph("Execute the tactical manual across a 4-week structured sprint:")

    table_d = doc.add_table(rows=5, cols=3)
    table_d.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_d.autofit = False
    
    headers_d = ["Execution Phase", "Primary Operational Mandate", "Verifiable Milestone Deliverable"]
    widths_d = [1.2, 2.8, 2.2]
    
    for i, title in enumerate(headers_d):
        cell = table_d.cell(0, i)
        cell.width = Inches(widths_d[i])
        set_cell_background(cell, "EDF2F7")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    gantt_data = [
        ("Week 1:\nBaseline & Bridge Launch", 
         "• Pull Employment Data Report from The Work Number; secure W-2s.\n• Contact one angel client; execute Chapter 2 SOW agreement.", 
         "[MILESTONE 1]: Verification Vault created. Bridge advisory project underway."),
         
        ("Week 2:\nDeliverable & ATS Engineering", 
         "• Deliver finished micro-consulting artifact (SOP or database).\n• Rebuild resume using Chapter 3 Single-Column Hybrid format; test plain-text parsing.", 
         "[MILESTONE 2]: Verifiable client deliverable archived. ATS-compliant resume finalized."),
         
        ("Week 3:\nToolchain & Script Conditioning", 
         "• Complete 72-hour AI workflow sprint (Notion/LLM structured prompt).\n• Record 30-second pivot scripts on phone voice memos until under 35 seconds.", 
         "[MILESTONE 3]: Modern portfolio asset stored. 30-second defense memorized."),
         
        ("Week 4:\nTargeted Outreach & Live Hunting", 
         "• Build 20-company watchlist (50–300 employees).\n• Send 5 weak-tie reconnection messages and 3 executive value-pitch letters.", 
         "[MILESTONE 4]: First 2 live informational interviews or screening calls scheduled.")
    ]
    
    for row_idx, data in enumerate(gantt_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table_d.cell(row_idx, col_idx)
            cell.width = Inches(widths_d[col_idx])
            set_cell_margins(cell, 80, 80, 100, 100)
            set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F7FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            else:
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

print("EN Appendices ready.")
