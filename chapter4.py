# -*- coding: utf-8 -*-
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import set_cell_background, set_cell_margins, add_callout
from chapter1 import add_heading_1, add_heading_2

def build_chapter_4(doc):
    add_heading_1(doc, "第四章：打破偏见——72小时现代 AI 与效率工具链极速武装")
    
    add_heading_2(doc, "1. 问题：断层带来的“老古董”标签")
    doc.add_paragraph(
        "即使求职者通过了简历初筛进入真人面试环节，依然会遭遇一道看不见但极度冰冷的审视壁垒："
        "面试官（尤其是比自己年轻 5 到 10 岁的用人部门组长或总监）会下意识地假设："
        "“这个人在家脱离职场整整两年，现在的日常办公早就普及了 AI 工具、自动化工作流和新一代协作系统，"
        "他肯定还停留在几年前的老旧工作模式里，招进来学习成本太高，思维肯定僵化生疏。”"
    )

    add_heading_2(doc, "2. 为什么发生：工具链爆炸制造的群体虚荣壁垒")
    doc.add_paragraph(
        "在 2023 至 2026 年间，生成式 AI（ChatGPT、Claude、Cursor、各主流大模型平台）以及现代协同工具（飞书多维表格、Notion AI、Zapier等）"
        "在企业内部完成了爆发式渗透。现职员工在日常工作中频繁使用这些工具，从而构建了一种排他性的“现代职场优越感”。"
        "他们认为只有每天浸泡在企业环境里的人才能掌握这些技术。"
    )
    doc.add_paragraph(
        "求职者之所以被这种偏见击中，是因为他们误以为“掌握现代工具”需要去读一个专业学位、或花费几个月从头学习编程。"
        "实际上，**当代绝大多数企业级 AI 和协同工具的设计理念就是“低代码/无代码与自然语言交互”**——"
        "一个具备 10 年深厚业务常识的老兵，只要掌握正确的提问与整合框架，花 72 小时就能在业务提效深度上，"
        "彻底秒杀那些只会用 AI 闲聊的初级在职员工。"
    )

    add_heading_2(doc, "3. 常见错误：陈旧套话与浮夸口号")
    doc.add_paragraph("在展示技能熟练度时，求职者最常犯的两类极端错误：")
    doc.add_paragraph("• 【极端一：停留在石器时代的名词】在简历上赫然写着“熟练掌握 Word、Excel、PPT、办公自动化、Outlook 邮件”。在当今的技术语境下，把这些当作核心技能写出来，等于主动在自己脸上盖上“我是计算机初级用户”的过时印章。")
    doc.add_paragraph("• 【极端二：脱离业务的假大空口号】在自我评价中宣称“精通 AI 大模型应用、具备 Prompt 提示词高阶能力”。当面试官随意追问一句“你在之前的财务分析中具体用了哪个模型的什么工作流时”，立刻语塞，暴露其只是浮于表面的概念搬运。")

    add_heading_2(doc, "4. 正确方法：业务场景绑定法（Tool-to-Workflow Binding）")
    doc.add_paragraph(
        "打破“脱节偏见”的破局法则，在于**【绝不孤立谈工具，只谈工具对具体业务动作的降本增效指标】**。\n"
        "我们不需要你成为大模型算法专家，只需要你在自己的专业领域（财务、运营、行政、法务、营销或技术）中，"
        "花 3 天时间，**亲手跑通一个用“现代协同工具 + 针对性大模型”完成的微型真实业务闭环**，并将其沉淀为一套可量化、可展示的案例话术。"
    )

    add_callout(
        doc,
        "面试官的心理定式与反制：\n"
        "面试官害怕的是“你思维僵化、抵触新生产力”；\n"
        "当你在面试中云淡风轻地聊出：“我习惯用 AI 辅助做复杂合同的跨版本 Diff 语义比对，通常能把初审效率提升 50%”时，"
        "这一句话产生的专业冲击力，会瞬间将面试官心中的“脱节老员工”预期，反转为“具备前沿实战降本能力的高手”。",
        title="偏见反转密码"
    )

    add_heading_2(doc, "5. 具体步骤：72 小时三日武装冲刺日程表")
    doc.add_paragraph("按照以下紧凑的三天计划，完成你对当代主流生产力工具的实战掌握：")

    doc.add_paragraph("【第 1 天：掌握现代结构化协同基座（以多维表格/Notion 为例）】\n"
                  "- 目标：抛弃简陋的本地传统 Excel，建立对“关系型多维表格与自动化看板”的直观理解。\n"
                  "- 动作：注册飞书（个人免费版）或 Notion，新建一个“多维表格（Database）”。练习如何建立单选标签、引用关联另一张表，并设置一条自动化触发器（例如：当某条任务状态改为‘已完成’时，自动发送通知）。耗时约 2～3 小时。")

    doc.add_paragraph("【第 2 天：攻克一个本专业的生成式 AI 结构化提效流】\n"
                  "- 目标：不再把 ChatGPT/Claude 当作玩具，而是当作一个严格受控的“初级分析师助手”。\n"
                  "- 动作：挑选你过去工作中做过的一份最复杂的业务文档（如一份老审计报告、一份活动方案、一份竞品对比），用标准结构化提示词（设定角色 + 提供背景 + 明确输出格式表格 + 给出约束条件）喂给大模型，练习让它在 3 分钟内输出高质量的结构化对比矩阵与数据清洗。耗时约 3 小时。")

    doc.add_paragraph("【第 3 天：凝练一份量化成果物并嵌入口径】\n"
                  "- 目标：把前两天的实战练习转化为简历上的一句硬核子弹陈述（Bullet Point），并在电脑里保存一份可随时展示的 PDF 成果截图。\n"
                  "- 动作：提炼标准句式：“熟练运用 [工具名称] 结合 [专业业务流]，实现 [具体指标提升]，具备自主搭建敏捷业务协同看板的能力”。")

    add_heading_2(doc, "6. 示例：四类常见职能岗位的“72小时改造对比表”")
    
    # 示例表格
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    headers = ["专业方向", "传统的陈旧表述（淘汰区）", "72小时武装后的现代实战表述（高分胜出区）"]
    col_widths = [1.2, 2.2, 2.8]
    
    for i, title in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, "EDF2F7")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    skill_data = [
        ("行政 / 综合运营", "“熟练使用 Office 办公软件，负责日常会议纪要、考勤统计与制度起草。”", "“熟练利用多维表格搭建团队敏捷考勤与资产台账看板；利用生成式大模型建立规章制度语义检索与智能初审流程，将常规政策解答效率提升60%。”"),
        ("财务 / 成本会计", "“精通 Excel 数据透视表与财务函数，能独立出具三大报表与成本核算。”", "“擅长通过 Python 基础脚本/现代自动化插件实现银行多流水与发票明细的自动化校验匹配，构建异常数据预警规则，大幅压缩月度结账对账周期。”"),
        ("内容 / 品牌营销", "“具备深厚文字功底，能独立撰写公关软文、双微运营与策划案。”", "“构建‘AI 辅助竞品监测与多平台内容资产分发工作流’，熟练利用结构化提示词进行多平台受众风格改写与社媒 SEO 关键词埋点，内容产出效率倍增。”"),
        ("项目管理 / PM", "“熟练使用 Project 软件，组织召开站会，跟进各部门项目进度。”", "“深度掌握 Jira / 飞书项目自动化流程，主导建立跨部门需求生命周期看板与甘特图联动机制；应用大模型自动归纳各端阻碍并生成周度风险简报。”")
    ]
    
    for row_idx, data in enumerate(skill_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, 80, 80, 100, 100)
            if row_idx % 2 == 1:
                set_cell_background(cell, "FFFFFF")
            else:
                set_cell_background(cell, "F7FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if col_idx == 2:
                run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            else:
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading_2(doc, "7. 检查清单：工具现代化 5 项过关自测")
    doc.add_paragraph("在面试中提及技术工具前，核对以下 5 条铁律：")
    doc.add_paragraph("【1】彻底清除了简历上所有“精通 Word/Excel/Windows系统”等严重暴露代际感的陈旧词汇。")
    doc.add_paragraph("【2】至少能流畅说出当前行业主流的 2 款数字化协作工具（如飞书、Notion、Slack、Jira 等）的核心优缺点。")
    doc.add_paragraph("【3】掌握大模型“结构化提问四要素”（角色、背景上下文、具体交付格式、排他性限制条件）。")
    doc.add_paragraph("【4】手头有一个亲自用 AI 或现代工具完成的脱敏成果物（如一张看板截图、一个报表清洗结果），可供随时佐证。")
    doc.add_paragraph("【5】谈及 AI 时，姿态始终是“业务总架构师在差遣数字化初级学徒”，绝不盲从 AI 的输出，强调人工校验与终审意识。")

    add_heading_2(doc, "8. 读者可以立即执行的行动")
    doc.add_paragraph(
        "【立即行动】：不要去看任何高深的编程网课。今天下午打开浏览器，在常用大模型（ChatGPT/Claude/国内主流平台）中，"
        "把自己过去最拿手的一份业务总结粘贴进去，输入：“*请扮演一家估值10亿的行业头部公司高管，用最严苛的眼光审视这份文档，"
        "指出其中 3 个未量化的逻辑漏洞，并输出一个改进后的三列对比表格*”。"
        "亲自体会一次现代高阶提示词的业务威力，你就完成了最关键的心智破冰。"
    )
    
    doc.add_page_break()

print("Chapter 4 module ready.")
