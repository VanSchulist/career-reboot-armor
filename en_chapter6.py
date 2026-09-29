# -*- coding: utf-8 -*-
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import set_cell_background, set_cell_margins
from make_english_book import add_heading_1, add_heading_2, add_callout

def build_en_chapter_6(doc):
    add_heading_1(doc, "Chapter 6: Bypassing the Portal Black Hole — Warm-Contact Scouting & Direct Value Pitches")
    
    add_heading_2(doc, "1. The Problem: The Exhaustion of Mass Portal Applications")
    doc.add_paragraph(
        "Even with an audited background history, a verifiable consulting bridge, an ATS-proof hybrid resume, and a rehearsed 30-second pivot, "
        "returners who rely exclusively on public job boards (LinkedIn 'Easy Apply,' Indeed, corporate career portals) face severe emotional attrition. "
        "Submitting 300 blind applications into corporate career sites yields an average interview callback rate below 1.5%. "
        "The relentless barrage of automated rejection emails slowly degrades a candidate's hard-won confidence, "
        "tempting them to surrender just when their tactical foundation is strongest."
    )

    add_heading_2(doc, "2. Why It Happens: The Open Market is an Algorithmic Meat Grinder")
    doc.add_paragraph(
        "Public job postings are the most crowded, commoditized tier of recruitment. A single posted role at a Fortune 500 company attracts "
        "between 500 and 2,000 applicants within 48 hours—including active employees recently laid off from high-profile tech or finance giants "
        "with continuous, unbroken resumes. In this high-volume environment, recruiters rely on brute-force keyword filters to instantly trim candidate pools."
    )
    doc.add_paragraph(
        "Attempting to compete solely on surface timeline metrics in a blind database is a statistical trap. Yet over 70% of professional positions "
        "are filled through the 'Hidden Job Market'—roles created through executive conversations, team expansions planned before public posting, "
        "or direct executive referrals. Returners must shift their battlefield to arenas where maturity, loyalty, and proven crisis resilience "
        "are recognized as premium assets."
    )

    add_heading_2(doc, "3. Common Mistakes: Desperation Begging vs. Cold LinkedIn Spam")
    doc.add_paragraph("Returners frequently squander their professional capital through two ineffective networking approaches:")
    doc.add_paragraph("• The Desperation Plea: Reaching out to old colleagues after three years of silence with an emotional message detailing their personal crisis, followed by 'Can you please submit my resume for any open role at your company?' This places an immense social burden on the recipient, resulting in awkward silence or polite evasion.")
    doc.add_paragraph("• Impersonal LinkedIn Spam: Sending copy-pasted InMails to陌生 hiring managers: 'Hi, I'm interested in your open role, here is my resume.' Recruiters delete these instantly because they fail to demonstrate any understanding of the company's specific operational challenges.")

    add_heading_2(doc, "4. The Correct Approach: The Warm-Contact Scouting Protocol")
    doc.add_paragraph(
        "The winning strategy is **TARGETED ASYMMETRIC OUTREACH**.\n"
        "Returners must redirect 80% of their job-search energy toward two receptive environments:\n"
        "1. Mid-Market Growth Companies (50 to 300 employees): These firms lack bureaucratic, hyper-rigid HR screening layers. Department Heads (VPs and Directors) make direct hiring decisions based on whether you can solve their immediate business bottlenecks, not arbitrary timeline algorithms;\n"
        "2. Dormant Weak-Tie Professional Networks: Engaging former colleagues and bosses through low-pressure **Informational Interviews**, seeking peer perspectives on industry shifts rather than begging for a job."
    )

    add_callout(
        doc,
        "THE NETWORKING PARADOX:\n"
        "If you ask someone for a job, you will receive advice and silence.\n"
        "If you ask someone for industry advice and their perspective on an operational challenge, you will often receive a job lead.\n"
        "Approaching peers as a competent equal preserves your executive presence and triggers referral reciprocity.",
        title="ASYMMETRIC RECRUITING RULE"
    )

    add_heading_2(doc, "5. Step-by-Step Implementation: The 4-Stage Warm-Scouting Pipeline")
    doc.add_paragraph("Execute this targeted outreach strategy to generate high-conversion live interviews:")

    doc.add_paragraph("Stage 1: Build a 20-Company Target Watchlist\n"
                  "Identify 20 profitable, mid-sized companies (50–300 employees) in your geographic area or target remote sector. "
                  "Criteria: Leadership has deep industry tenure; company recently raised Series B/C funding or expanded product lines; hiring decisions rest with functional VPs rather than third-party RPO recruiters.")

    doc.add_paragraph("Stage 2: Re-activate Dormant Weak Ties with Low-Friction Inquiries\n"
                  "Identify 5 former respected colleagues or managers. Send a short, pressure-free re-connection message:\n"
                  "'Hi [Name]! Hope you're doing great. After a planned sabbatical managing family medical affairs (which are now fully and happily resolved), I'm actively stepping back into the industry. I've been researching [Specific Modern Trend, e.g., AI integration in logistics] and saw your team's impressive work at [Company]. I know you're slammed, but I'd love to buy you a 15-minute virtual coffee sometime next week just to get your pulse on how the landscape has shifted over the last couple of years. Zero expectations—just great to reconnect!'")

    doc.add_paragraph("Stage 3: Conduct the 15-Minute Peer Intelligence Chat\n"
                  "During the conversation, do NOT ask for a referral. Focus entirely on listening to their team's current friction points. "
                  "Casually reference the modern tools you've deployed during your consulting bridge. In 65% of cases, when they realize you are sharp, unencumbered, and available, they will ask: 'Are you looking right now? We actually have a headcount opening that hasn't been posted yet.'")

    doc.add_paragraph("Stage 4: Deploy the Executive Value Pitch Letter directly to Department Heads\n"
                  "For target companies with unadvertised needs, email the functional VP directly (finding their email via Apollo.io or Hunter.io). "
                  "Do not send a generic cover letter. Send a crisp, 3-paragraph Value Pitch diagnosing a specific problem they face, backed by your 1-page hybrid resume.")

    add_heading_2(doc, "6. Concrete Example: The Cold Executive Value Pitch Letter")
    
    # Sample Letter Box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, 120, 120, 140, 140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    
    sample_pitch = (
        "SUBJECT: Observations on scaling regional distribution — and an introduction\n\n"
        "Dear Mr. Vance [VP of Supply Chain / Operations],\n\n"
        "I've been following Apex Logistics' recent expansion into mid-Atlantic cold-storage fulfillment with great interest.\n\n"
        "Having previously served as Director of Operations at Global Logistics Corp—where I led multi-site distribution hubs handling $40M+ in volume—"
        "I know that rapid regional expansions frequently encounter 10–15% margin leakage through WMS inventory discrepancies and carrier reconciliation delays. "
        "Recently, while consulting for a mid-market distributor, I designed an automated variance-tracking protocol that cut inventory reconciliation time from 7 days to 48 hours.\n\n"
        "Following a planned family sabbatical that has now fully and permanently concluded, I am returning to full-time executive operations. "
        "I've attached my single-page hybrid resume for your reference. If Apex is currently navigating fulfillment bottlenecks and could benefit from "
        "a seasoned operator who can hit the ground running without ramp-up friction, I would welcome a brief 15-minute introductory conversation.\n\n"
        "Regardless, congratulations on the continued regional expansion.\n\n"
        "Best regards,\n\n"
        "David K. Miller | (215) 555-0192 | linkedin.com/in/david-miller-ops | [Attached: David_Miller_Resume.pdf]"
    )
    run_p = p.add_run(sample_pitch)
    run_p.font.name = 'Consolas'
    run_p.font.size = Pt(8.5)
    run_p.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading_2(doc, "7. Practical Checklist: The Outreach Discipline Audit")
    doc.add_paragraph("In managing your outreach campaign, verify:")
    doc.add_paragraph("[1] At least 60% of applications target mid-sized firms (50–300 employees) rather than monolithic enterprise portals.")
    doc.add_paragraph("[2] Weak-tie re-engagement messages contain zero desperate pleas or explicit job requests.")
    doc.add_paragraph("[3] Executive Value Pitch letters lead with the employer's business problem, not your personal need for employment.")
    doc.add_paragraph("[4] Identified the direct functional department head (VP/Director) rather than sending unsolicited PDFs to general HR inboxes.")
    doc.add_paragraph("[5] Maintained an organized outreach tracker (Company, Contact, Date Sent, Follow-up Date).")
    doc.add_paragraph("[6] Mentally prepared for a 40–50% non-response rate as normal commercial reality, maintaining zero self-worth erosion.")

    add_heading_2(doc, "8. Immediate Actionable Step")
    doc.add_paragraph(
        "IMMEDIATE ACTION TODAY: Open LinkedIn. Identify 3 former colleagues or managers who respect your historical work and currently work "
        "at stable organizations. Send the low-pressure Stage 2 reconnection message to the first contact. "
        "You have officially taken control of your professional pipeline."
    )
    
    doc.add_page_break()

print("EN Chapter 6 ready.")
