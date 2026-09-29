# -*- coding: utf-8 -*-
"""
Master Assembly Script for 'The Career Reboot Armor: US Edition'
Compiles Preface, Chapters 1-6, Appendices A-D, and Epilogue into a production Word (.docx) document.
"""

import os
import docx
from make_english_book import create_english_document, add_heading_1, add_heading_2, add_callout
from en_chapter1 import build_en_chapter_1
from en_chapter2 import build_en_chapter_2
from en_chapter3 import build_en_chapter_3
from en_chapter4 import build_en_chapter_4
from en_chapter5 import build_en_chapter_5
from en_chapter6 import build_en_chapter_6
from en_appendices import build_en_appendices
from docx.shared import Pt, RGBColor

def build_en_preface(doc):
    add_heading_1(doc, "Preface: An Operational Field Order for Resilient Returners")
    doc.add_paragraph(
        "This tactical manual is written exclusively for mid-to-senior professionals who have endured major human trials—navigating "
        "the agonizing terminal illness of an aging parent, supporting a critically ill spouse or child through multi-year medical crises, "
        "or surviving a serious personal medical battle—and are now stepping forward to reclaim their careers in the United States corporate arena."
    )
    doc.add_paragraph(
        "When confronted with domestic upheaval and life-or-death realities, you made the honorable, responsible choice to stand your post as "
        "a caregiver or patient. But when you prepare to re-enter the workforce, the contemporary American recruitment machinery greets you with "
        "calculated indifference: Applicant Tracking Systems (ATS) flag multi-year gaps as critical disqualifiers; corporate recruiters reflexively "
        "discard resumes with inactive dates; and third-party background screening companies cast an ominous shadow over candidates who fear an unbridgeable void."
    )
    doc.add_paragraph(
        "We reject toxic positivity, sentimental encouragement, and patronizing platitudes. We will never advise you to describe your caregiving as "
        "'Chief Household Officer' or claim that 'managing domestic budgets makes you an enterprise CFO.' Corporate recruitment operates on cold, "
        "risk-adjusted economic utility. The purpose of this guide is to equip you with an airtight operational armor: legally sound verification "
        "tactics, algorithmic resume engineering, verified consulting bridges, defensive interview scripting, and targeted direct-outreach playbooks."
    )
    add_callout(
        doc,
        "CORE RULES OF ENGAGEMENT:\n"
        "1. Sequential Execution: Follow Chapters 1 through 6 in exact order. First, audit your background check baseline (FCRA/The Work Number); "
        "second, establish a verifiable micro-consulting bridge; third, re-engineer your resume into an ATS-proof single-column hybrid; fourth, sprint "
        "through the modern AI/collaboration toolchain; fifth, rehearse your 30-second interview firewalls; and sixth, launch direct outreach to mid-market decision-makers.\n"
        "2. Strict Legal Truthfulness: Never alter employment dates or invent fictitious W-2 employers. Modern background checks verify exact dates via automated payroll feeds. "
        "Tactical defense relies on legally protected gaps and verified 1099/advisory engagements, not fabrication.\n"
        "3. Dignity is Built on Competence: Respect in corporate America is not granted out of sympathy; it is claimed through crisp professionalism, modern tool mastery, and unshakeable composure.",
        title="RULES OF ENGAGEMENT"
    )
    doc.add_page_break()

def build_en_epilogue(doc):
    doc.add_page_break()
    add_heading_1(doc, "Epilogue: Return with Scars and Medals")
    doc.add_paragraph(
        "When you have completed the protocols outlined in this field manual—locking down your employment verification baseline, delivering verifiable "
        "micro-consulting artifacts, structuring your single-column hybrid resume, mastering modern AI workflows, delivering flawless 30-second defensive pivots, "
        "and initiating value-first outreach to business leaders—you are no longer the hesitant candidate staring at an intimidating career gap in the dead of night."
    )
    doc.add_paragraph(
        "Always remember this fundamental truth: Standing by an ailing parent during their final years, fiercely protecting a medically fragile loved one, "
        "or fighting your way back from severe personal illness is one of the highest expressions of moral character, courage, and stamina a human being can demonstrate. "
        "It is a medal of honor, not a professional stain."
    )
    doc.add_paragraph(
        "Those who spent those same years in climate-controlled offices may possess continuous, uninterrupted timelines. But candidates forged in the crucible "
        "of real-world adversity possess executive capabilities that cannot be taught in business schools: unflinching emotional composure under acute pressure, "
        "the discernment to separate real crises from routine office friction, and an unshakeable perspective on what truly matters."
    )
    doc.add_paragraph(
        "Discard the guilt. Don your tactical armor. Step confidently back into the arena—your place at the table is waiting."
    )

def main():
    print("=" * 60)
    print("Starting compilation of 'The Career Reboot Armor: US Edition'...")
    print("=" * 60)
    
    print("[1/9] Initializing Word document and global typography styles...")
    doc = create_english_document()
    
    print("[2/9] Generating Preface...")
    build_en_preface(doc)
    
    print("[3/9] Compiling Chapter 1: Demystifying Background Checks & FCRA...")
    build_en_chapter_1(doc)
    
    print("[4/9] Compiling Chapter 2: The Consulting Umbrella & SOW Bridges...")
    build_en_chapter_2(doc)
    
    print("[5/9] Compiling Chapter 3: ATS-Proof Hybrid Chronological Resume...")
    build_en_chapter_3(doc)
    
    print("[6/9] Compiling Chapter 4: Crushing Stale-Skills Bias with AI Sprints...")
    build_en_chapter_4(doc)
    
    print("[7/9] Compiling Chapter 5: 30-Second Pivot & Interview Privacy Firewalls...")
    build_en_chapter_5(doc)
    
    print("[8/9] Compiling Chapter 6: Bypassing Job Boards via Mid-Market Outreach...")
    build_en_chapter_6(doc)
    
    print("[9/9] Compiling Complete Production Appendices (A, B, C, D)...")
    build_en_appendices(doc)
    
    print("Adding Epilogue...")
    build_en_epilogue(doc)
    
    target_path = os.path.join(os.getcwd(), "The_Career_Reboot_Armor_US_Edition.docx")
    print(f"Saving compiled document to: {target_path}...")
    doc.save(target_path)
    
    file_size_kb = os.path.getsize(target_path) / 1024
    print(f"Document compilation complete! Final size: {file_size_kb:.2f} KB")
    print("=" * 60)

if __name__ == "__main__":
    main()
