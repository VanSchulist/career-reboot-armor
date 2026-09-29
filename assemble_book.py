# -*- coding: utf-8 -*-
import os
from make_book import create_document
from chapter1 import build_chapter_1, add_heading_1, add_heading_2
from chapter2 import build_chapter_2
from chapter3 import build_chapter_3
from chapter4 import build_chapter_4
from chapter5 import build_chapter_5
from chapter6 import build_chapter_6
from appendices import build_appendices
from docx_helpers import add_callout
from docx.shared import Pt, RGBColor

def build_preface(doc):
    add_heading_1(doc, "前言：写给逆境归来者的作战序言")
    doc.add_paragraph(
        "这是一本专为经历过重大人间冷暖、承担过沉重家庭伦理责任，如今准备重披战袍的专业人士编写的求职实战战术手册。"
    )
    doc.add_paragraph(
        "在过往的几年里，你或许在医院长廊彻夜守候过病危的父母，或许在特护病房里陪伴过至亲走完最后一程，"
        "或许经历了重大疾病的生死涅槃。面对家庭与生命的剧烈风暴，你做出了身为儿女或亲人最负责任的崇高选择。"
        "然而，当你历经千辛万苦想要重返职场时，现代商业招聘体系却以极其冷酷的姿态迎接了你——"
        "ATS 算法将你的空白期判定为缺陷，第三方背调的阴影让你提心吊胆，面试官带着怀疑的目光审视你是否脱节。"
    )
    doc.add_paragraph(
        "我们拒绝任何虚浮的励志毒鸡汤，也不劝你把家务伪装成管理经验去自取其辱。"
        "商业世界遵循冰冷的交换法则。本书的目的，是为你锻造一套经得起背景调查、穿透机器算法、抵御探私盘问的【硬核战略护甲】。"
    )
    add_callout(
        doc,
        "全书阅读与执行守则：\n"
        "1. 本书不讲抽象大道理，每一章均解决一个致命断层，附带具体的执行步骤、对照表格与行动清单。\n"
        "2. 请严格按照从第一章到第六章的顺序推进：先清算背调事实，再搭建过桥项目，接着重构简历版式，极速武装现代工具，背熟面试剧本，最后定向进攻腰部企业。\n"
        "3. 尊严不是别人施舍的，而是用严密的专业素养和滴水不漏的战术自己赢回来的。",
        title="本书使用守则"
    )
    doc.add_page_break()

def build_epilogue(doc):
    add_heading_1(doc, "结语：带着伤疤与勋章，重返你的战场")
    doc.add_paragraph(
        "当你按照本书的步骤，完成个人数据的基准锁定、交付了真实的过桥项目、重构了单栏混合型简历、"
        "熟练掌握了 30 秒面试防线并发出第一封温差自荐信时，你已经彻底脱离了那个在深夜里对着空白简历自卑焦虑的自己。"
    )
    doc.add_paragraph(
        "请永远记住：**在人生的长跑中，照顾病重垂危的父母、在生死关头挺身而出，是一个人道德人格最辉煌的勋章，绝不是你职业生涯的污点。**"
    )
    doc.add_paragraph(
        "那些在平坦温室中从未经历过家庭变故的人，或许拥有未曾中断的完美时间线；"
        "但经历过这场极限淬炼的你，拥有他们未曾拥有的顶级核心竞争力——在绝境中的情绪稳定性、在巨大压力下的多线斡旋力、以及对生活本质的深刻洞察。"
    )
    doc.add_paragraph(
        "收起心虚，穿好这套护甲，昂首挺胸走回属于你的战场。"
    )

def main():
    print("Initializing document...")
    doc = create_document()
    
    print("Building Preface...")
    build_preface(doc)
    
    print("Building Chapter 1...")
    build_chapter_1(doc)
    
    print("Building Chapter 2...")
    build_chapter_2(doc)
    
    print("Building Chapter 3...")
    build_chapter_3(doc)
    
    print("Building Chapter 4...")
    build_chapter_4(doc)
    
    print("Building Chapter 5...")
    build_chapter_5(doc)
    
    print("Building Chapter 6...")
    build_chapter_6(doc)
    
    print("Building Appendices...")
    build_appendices(doc)
    
    print("Building Epilogue...")
    build_epilogue(doc)
    
    target_path = os.path.join(os.getcwd(), "重返职场战略护甲_求职实战指南.docx")
    print(f"Saving to {target_path}...")
    doc.save(target_path)
    print("Word document built successfully!")

if __name__ == "__main__":
    main()
