# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx_helpers import set_cell_background, set_cell_margins, add_callout

def create_document():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles setup
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Microsoft YaHei'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    style_normal.paragraph_format.line_spacing = 1.35
    style_normal.paragraph_format.space_after = Pt(6)
    
    # Title Page
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("重返职场战略护甲\n长病假与长期照护后的履历过桥与求职实战指南")
    run_title.font.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(28)
    run_sub = p_sub.add_run("The Career Reboot Armor: A Field Guide to Gap Bridging & Job Search")
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    
    # Metadata Box
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(40)
    run_meta = p_meta.add_run("适用人群：因照护失能/患病亲人或重大伤病康复，离开职场1～4年的重返者\n编写原则：基于真实招聘机制与背景调查规则，拒绝虚假情怀与违法造假")
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    doc.add_page_break()
    return doc

print("Init module ready.")
